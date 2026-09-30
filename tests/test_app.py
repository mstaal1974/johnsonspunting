import io
from datetime import date, datetime

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.db import SessionLocal
from app.main import app
from app.models import Bet
from app.services.club import build_view
from app.services.excel_io import import_workbook, split_bets


def club_workbook() -> bytes:
    """Minimal copy of the club spreadsheet layout."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bets"
    for m in range(12):
        c = 3 + m * 4
        y, mo = divmod(9 + m, 12)
        ws.cell(1, c, datetime(2025 + y, mo + 1, 1))
        for i, h in enumerate(["Stake to bet this month", "Bet Placed Y/N", "Bets", "Collect this month"]):
            ws.cell(2, c + i, h)
    rows = [
        ("Greg", "$30 NRL ATT X Coates @ $1.93; $30 NRLW ATT J Fressard @ $2.25; $30 NRL/W SGM 3 Legs @ $7.74;", 0),
        ("Brad ", "$20W Antino @ $8.50; $20 BONER Bx Tri @ 83%; $30 GGs Multi 2 Legs @ $12.17;", 0),
        ("Tom", "$25W Tentyris @ $6.50; $25W Lazzura @ $6.50;", 162.5),
        ("Col", None, 0),
    ]
    for r, (name, text, collect) in enumerate(rows, start=3):
        ws.cell(r, 1, name)
        ws.cell(r, 3, 50)
        ws.cell(r, 5, text)
        ws.cell(r, 6, collect)
    ws.cell(3 + len(rows), 1, "TOTAL COLLECT")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_split_bets():
    parsed = split_bets("$20W Antino @ $8.50; $20 BONER Bx Tri @ 83%;")
    assert parsed == [(20.0, "$20W Antino @ $8.50", "$8.50"), (20.0, "$20 BONER Bx Tri @ 83%", "83%")]


def test_import_reproduces_spreadsheet_rules():
    with TestClient(app):  # runs startup (creates tables + seed)
        pass
    with SessionLocal() as db:
        res = import_workbook(db, club_workbook(), today=date(2025, 11, 10))
        assert res["bets"] == 8 and not res["warnings"]
        view = build_view(db, today=date(2025, 11, 10))
        by = {m.name: view.summaries[m.id] for m in view.members}
        assert by["Greg"].months[1].available == 10       # overspent $40 in Oct
        assert by["Brad"].months[1].available == 50       # bonus bet doesn't count
        assert by["Tom"].banked == 81.25                   # half of $162.50
        assert by["Tom"].months[1].available == 131.25     # $50 + other half
        assert by["Tom"].brownlow == 3
        assert by["Col"].months[0].unused == 50
        assert view.ranked()[0][1].name == "Tom"


def test_admin_flow():
    with TestClient(app) as c:
        assert c.get("/").status_code == 200
        r = c.get("/admin", follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/login"
        assert "Wrong password" in c.post("/login", data={"password": "nope"}).text
        c.post("/login", data={"password": "test-pw"})

        with SessionLocal() as db:
            view = build_view(db)
            month = view.default_month
            ben = next(m for m in view.members if m.name == "Ben")
        r = c.post("/admin/bets", data={"month": month, "member_id": ben.id, "stake": "$30",
                                        "description": "W Snitzel Dancer", "odds": "$4.80"})
        assert r.status_code == 200 and "Saved Ben" in r.text
        with SessionLocal() as db:
            bet = db.query(Bet).filter(Bet.member_id == ben.id).order_by(Bet.id.desc()).first()
            assert bet.result == "pending" and bet.stake == 30

        r = c.post(f"/admin/bets/{bet.id}/settle", data={"result": "won", "collect": "144"})
        assert "$72.00 banked" in r.text
        board = c.get("/api/leaderboard").json()["leaderboard"]
        assert next(row for row in board if row["punter"] == "Ben")["banked"] == 72

        r = c.post("/admin/bets", data={"month": month, "member_id": ben.id, "stake": "abc"})
        assert "Stake must be a number" in r.text

        for path in ("/", "/?sort=brownlow", f"/month/{month}", f"/punter/{ben.id}", "/admin",
                     f"/admin/bets/{bet.id}", "/admin/members", "/admin/data"):
            assert c.get(path).status_code == 200, path

        xlsx = c.get("/export.xlsx")
        wb = load_workbook(io.BytesIO(xlsx.content))
        assert wb.sheetnames == ["Leaderboard", "Bets", "Total Banked", "Brownlow"]
        assert "Ben" in [c.value for c in wb["Leaderboard"]["B"]]

        c.post(f"/admin/bets/{bet.id}/delete")
        with SessionLocal() as db:
            assert db.get(Bet, bet.id) is None


def test_favicon():
    with TestClient(app) as c:
        r = c.get("/favicon.ico")
        assert r.status_code == 200 and r.headers["content-type"].startswith("image/svg+xml")


def test_database_failure_shows_reason(monkeypatch):
    from app import db

    def boom():
        raise RuntimeError("connection refused")

    monkeypatch.setattr(db, "init_db", boom)
    r = TestClient(app).get("/")
    assert r.status_code == 503 and "connection refused" in r.text
