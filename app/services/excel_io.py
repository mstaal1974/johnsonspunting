"""Import from / export to the club's existing spreadsheet layout.

The 'Bets' sheet has punter names in column A (from row 3) and a block of four
columns per month starting at column C:
    Stake to bet this month | Bet Placed Y/N | Bets | Collect this month
Row 1 holds the month's date above each block. 'Bets' is free text with one
bet per ';' e.g. "$20W Antino @ $8.50; $20 BONER Bx Tri @ 83%;".
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import engine
from app.models import Bet, Member, Season
from app.services.club import SeasonView, season_name

STAKE_RE = re.compile(r"^\$?\s*(\d+(?:\.\d+)?)")
BONUS_RE = re.compile(r"\bboner\b|bonus", re.I)
ODDS_RE = re.compile(r"@\s*(\$?\d+(?:\.\d+)?%?(?:/\$?\d+(?:\.\d+)?)?)")


@dataclass
class ParsedBet:
    name: str
    month: int
    stake: float
    description: str
    odds: str | None
    bonus: bool = False


@dataclass
class ParsedSheet:
    start: date
    names: list[str]
    bets: list[ParsedBet] = field(default_factory=list)
    collects: dict[tuple[str, int], float] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def split_bets(text: str) -> list[tuple[float | None, str, str | None]]:
    out = []
    for part in str(text).split(";"):
        part = part.strip()
        if not part:
            continue
        m = STAKE_RE.match(part)
        o = ODDS_RE.search(part)
        out.append((float(m.group(1)) if m else None, part, o.group(1) if o else None))
    return out


def _as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def parse_workbook(data: bytes) -> ParsedSheet:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    if "Bets" not in wb.sheetnames:
        raise ValueError("Workbook has no 'Bets' sheet")
    ws = wb["Bets"]

    month_cols: list[tuple[int, date]] = []
    col = 3
    while col <= ws.max_column and len(month_cols) < engine.MONTHS_PER_SEASON:
        d = _as_date(ws.cell(1, col).value)
        if d is None:
            break
        month_cols.append((col, d))
        col += 4
    if not month_cols:
        raise ValueError("Couldn't find month dates in row 1 of the 'Bets' sheet")
    start = month_cols[0][1].replace(day=1)

    parsed = ParsedSheet(start=start, names=[])
    for row in range(3, ws.max_row + 1):
        raw = ws.cell(row, 1).value
        if raw is None or not str(raw).strip():
            break
        name = str(raw).strip()
        if name.upper().startswith("TOTAL"):
            break
        parsed.names.append(name)
        for m, (c, _) in enumerate(month_cols):
            text = ws.cell(row, c + 2).value
            collect = ws.cell(row, c + 3).value
            if isinstance(collect, (int, float)) and collect > 0:
                parsed.collects[(name, m)] = float(collect)
            if not text or not str(text).strip():
                if (name, m) in parsed.collects:
                    parsed.warnings.append(f"{name} {month_cols[m][1]:%b}: collect but no bets listed")
                continue
            for stake, desc, odds in split_bets(text):
                if stake is None:
                    parsed.warnings.append(f"{name} {month_cols[m][1]:%b}: no stake found in '{desc}' - imported as $0")
                parsed.bets.append(ParsedBet(name, m, stake or 0.0, desc, odds, bool(BONUS_RE.search(desc))))
    return parsed


def import_workbook(db: Session, data: bytes, today: date | None = None, zero_collect_is_lost: bool = True) -> dict:
    """Load the spreadsheet into the database, replacing that season's bets.

    Collects are recorded per month in the sheet, not per bet, so a month's
    collect is put on its first bet and the month's other bets are marked lost.
    Bets in finished months with no collect are marked lost when
    `zero_collect_is_lost`; bets in the current month stay pending.
    """
    today = today or date.today()
    p = parse_workbook(data)

    season = db.query(Season).filter(Season.start == p.start).first()
    if season is None:
        season = Season(name=season_name(p.start), start=p.start, base_stake=50.0, max_bets=2)
        db.add(season)
    for s in db.query(Season).all():
        s.active = False
    season.active = True
    db.flush()

    members = {m.name.lower(): m for m in db.query(Member).all()}
    for name in p.names:
        if name.lower() not in members:
            m = Member(name=name)
            db.add(m)
            members[name.lower()] = m
    db.flush()

    db.query(Bet).filter(Bet.season_id == season.id).delete()
    current = engine.current_month_index(p.start, today)
    first_bet_in_month: set[tuple[str, int]] = set()
    for pb in p.bets:
        key = (pb.name, pb.month)
        collect = p.collects.get(key, 0.0)
        if collect > 0:
            result = engine.WON if key not in first_bet_in_month else engine.LOST
            bet_collect = collect if result == engine.WON else 0.0
        elif pb.month < current and zero_collect_is_lost:
            result, bet_collect = engine.LOST, 0.0
        else:
            result, bet_collect = engine.PENDING, 0.0
        first_bet_in_month.add(key)
        db.add(Bet(
            season_id=season.id, member_id=members[pb.name.lower()].id, month=pb.month,
            stake=pb.stake, description=pb.description, odds=pb.odds,
            result=result, collect=bet_collect, bonus=pb.bonus,
        ))
    db.commit()
    return {"season": season.name, "members": len(p.names), "bets": len(p.bets), "warnings": p.warnings}


HEADER_FILL = PatternFill("solid", fgColor="1F4E3D")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _header(ws, row: int, values: list) -> None:
    for i, v in enumerate(values, start=1):
        c = ws.cell(row, i, v)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")


def export_workbook(view: SeasonView) -> bytes:
    wb = Workbook()
    money = '"$"#,##0.00'

    ws = wb.active
    ws.title = "Leaderboard"
    ws["A1"] = f"Johnsons Punt Club {view.season.name} - {view.current_label}"
    ws["A1"].font = Font(bold=True, size=14)
    _header(ws, 3, ["Rank", "Punter", "Total Banked", "Contributed", "Net", "ROI",
                    "Total Collected", "Total Staked", "Bets", "Winners", "Brownlow"])
    for r, (rank, m, s) in enumerate(view.ranked(), start=4):
        ws.append([rank, m.name, s.banked, s.contributed, s.net, s.roi, s.collect, s.staked,
                   s.bets, s.winners, s.brownlow])
        for col in (3, 4, 5, 7, 8):
            ws.cell(r, col).number_format = money
        ws.cell(r, 6).number_format = "0.0%"
    ws.column_dimensions["B"].width = 14
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 14

    ws = wb.create_sheet("Bets")
    ws.cell(2, 1, "Punter")
    ws.cell(2, 2, "Total Collected")
    for m, label in enumerate(view.labels):
        c = 3 + m * 4
        ws.cell(1, c, label).font = Font(bold=True)
        ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + 3)
        for i, h in enumerate(["Stake to bet this month", "Bet Placed Y/N", "Bets", "Collect this month"]):
            ws.cell(2, c + i, h)
        ws.column_dimensions[get_column_letter(c + 2)].width = 40
    for c in range(1, 3 + 4 * len(view.labels)):
        ws.cell(2, c).fill, ws.cell(2, c).font = HEADER_FILL, HEADER_FONT
        ws.cell(2, c).alignment = Alignment(wrap_text=True)
    bets_by = {}
    for b in view.bets:
        bets_by.setdefault((b.member_id, b.month), []).append(b)
    for r, m in enumerate(view.members, start=3):
        s = view.summaries[m.id]
        ws.cell(r, 1, m.name)
        ws.cell(r, 2, s.collect).number_format = money
        for line in s.months:
            c = 3 + line.month * 4
            mb = bets_by.get((m.id, line.month), [])
            ws.cell(r, c, line.available).number_format = money
            ws.cell(r, c + 1, "y" if mb else "")
            ws.cell(r, c + 2, "; ".join(b.description for b in mb))
            ws.cell(r, c + 3, line.collect).number_format = money

    for title, attr in (("Total Banked", "banked"), ("Brownlow", "brownlow")):
        ws = wb.create_sheet(title)
        _header(ws, 2, ["Name", title if attr == "banked" else "Total"] + [l[:3] for l in view.labels])
        for r, m in enumerate(view.members, start=3):
            s = view.summaries[m.id]
            ws.cell(r, 1, m.name)
            ws.cell(r, 2, getattr(s, attr))
            for line in s.months:
                ws.cell(r, 3 + line.month, getattr(line, attr))
            if attr == "banked":
                for c in range(2, 3 + len(view.labels)):
                    ws.cell(r, c).number_format = money

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
